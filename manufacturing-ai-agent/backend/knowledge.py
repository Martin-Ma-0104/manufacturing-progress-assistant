"""Local extraction + persistent Chroma index. No automatic model downloads/egress."""

from collections import Counter
from contextlib import closing
import hashlib
import io
import math
import os
from pathlib import Path
import re
import sqlite3
import zipfile

os.environ.setdefault("ANONYMIZED_TELEMETRY", "False")
DIM = 384
CATEGORIES = {
    "Production SOP",
    "Quality Standards",
    "Equipment Manuals",
    "Maintenance Procedures",
    "Troubleshooting Guidelines",
    "Product Specifications",
    "Manufacturing Process Knowledge",
}


def storage():
    return Path(
        os.getenv("KNOWLEDGE_DIR", str(Path(__file__).parent / "storage"))
    ).resolve()


def connect():
    storage().mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(storage() / "knowledge.sqlite3")
    db.execute(
        "CREATE TABLE IF NOT EXISTS documents (id TEXT PRIMARY KEY,name TEXT,category TEXT,provider TEXT,text TEXT,pages INTEGER)"
    )
    db.execute(
        "CREATE TABLE IF NOT EXISTS chunks (id TEXT PRIMARY KEY,document_id TEXT,location TEXT,text TEXT)"
    )
    return db


def provider():
    return os.getenv("EMBEDDING_PROVIDER", "local-hash-v1")


def index_identity():
    # Vector spaces cannot be reused across embedding model configurations.
    return (
        "openai:" + os.getenv("EMBEDDING_MODEL", "text-embedding-3-small")
        if provider() == "openai"
        else provider()
    )


def embeddings(texts, allow_external=False):
    if provider() == "openai":
        if not allow_external:
            raise ValueError(
                "External embedding requires explicit permission for document/query text."
            )
        from openai import OpenAI

        key = os.getenv("LLM_API_KEY") or os.getenv("OPENAI_API_KEY")
        if not key:
            raise ValueError("Embedding API is not configured.")
        with OpenAI(api_key=key, max_retries=0, timeout=20) as client:
            response = client.embeddings.create(
                model=os.getenv("EMBEDDING_MODEL", "text-embedding-3-small"),
                input=texts,
            )
            return [r.embedding for r in response.data]
    if provider() != "local-hash-v1":
        raise ValueError("Unsupported embedding provider.")
    vectors = []
    for text in texts:
        text = re.sub(r"\s+", " ", text.lower())
        features = re.findall(r"\w+", text) + [
            text[i : i + 3] for i in range(max(0, len(text) - 2))
        ]
        vector = [0.0] * DIM
        for term, count in Counter(features).items():
            digest = hashlib.blake2b(term.encode(), digest_size=8).digest()
            vector[int.from_bytes(digest, "big") % DIM] += 1 + math.log(count)
        norm = math.sqrt(sum(v * v for v in vector)) or 1
        vectors.append([v / norm for v in vector])
    return vectors


def collection():
    import chromadb
    from chromadb.config import Settings

    client = chromadb.PersistentClient(
        path=str(storage() / "vectors"), settings=Settings(anonymized_telemetry=False)
    )
    return client.get_or_create_collection(
        "separator_" + hashlib.sha256(index_identity().encode()).hexdigest()[:16],
        embedding_function=None,
        metadata={"hnsw:space": "cosine"},
    )


def extract(name, data):
    ext = Path(name).suffix.lower()
    if not data or len(data) > 5 * 1024 * 1024:
        raise ValueError("Upload a nonempty file up to 5 MiB.")
    if ext in (".txt", ".md"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeError:
            raise ValueError("Text documents must use UTF-8.") from None
        pages = [("Text", text)]
    elif ext == ".pdf":
        from pypdf import PdfReader

        if not data.startswith(b"%PDF-"):
            raise ValueError("Invalid PDF signature.")
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted or len(reader.pages) > 200:
            raise ValueError("Use an unencrypted PDF with at most 200 pages.")
        pages = [
            (f"Page {i + 1}", p.extract_text() or "")
            for i, p in enumerate(reader.pages)
        ]
    elif ext == ".docx":
        from docx import Document

        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > 20 * 1024 * 1024:
                raise ValueError("Expanded DOCX exceeds the size limit.")
        doc = Document(io.BytesIO(data))
        pages = [(f"Paragraph {i + 1}", p.text) for i, p in enumerate(doc.paragraphs)]
        pages += [
            (
                f"Table {i + 1}",
                "\n".join(" | ".join(c.text for c in row.cells) for row in t.rows),
            )
            for i, t in enumerate(doc.tables)
        ]
    else:
        raise ValueError(
            "Supported extensions: PDF, DOCX, TXT, MD. No OCR or executable files."
        )
    if sum(len(t) for _, t in pages) > 250000:
        raise ValueError("Extracted text exceeds 250,000 characters.")
    if not any(t.strip() for _, t in pages):
        raise ValueError("No extractable text; image-only documents need separate OCR.")
    return pages


def ingest(name, data, category="Production SOP", allow_external=False):
    if category not in CATEGORIES:
        raise ValueError("Unknown document category.")
    pages = extract(name, data)
    doc_id = hashlib.sha256(data).hexdigest()[:24]
    with closing(connect()) as db:
        existing = db.execute(
            "SELECT provider FROM documents WHERE id=?", (doc_id,)
        ).fetchone()
        if existing:
            if existing[0] != index_identity():
                raise ValueError(
                    "Embedding model changed. Use a separate KNOWLEDGE_DIR to rebuild the index; incompatible vector spaces cannot be mixed."
                )
            return dict(id=doc_id, duplicate=True)
        chunks = []
        for location, text in pages:
            for offset in range(0, len(text), 700):
                part = text[offset : offset + 850].strip()
                if part:
                    chunks.append((f"{doc_id}-{len(chunks)}", doc_id, location, part))
        if len(chunks) > 400:
            raise ValueError("Too many chunks; split the document.")
        vectors = embeddings([c[3] for c in chunks], allow_external)
        col = collection()
        col.upsert(
            ids=[c[0] for c in chunks],
            embeddings=vectors,
            metadatas=[{"document_id": doc_id} for c in chunks],
        )
        try:
            with db:
                db.execute(
                    "INSERT INTO documents VALUES (?,?,?,?,?,?)",
                    (
                        doc_id,
                        Path(name).name[:200],
                        category,
                        index_identity(),
                        "\n".join(t for _, t in pages),
                        len(pages),
                    ),
                )
                db.executemany("INSERT INTO chunks VALUES (?,?,?,?)", chunks)
        except Exception:
            col.delete(ids=[c[0] for c in chunks])
            raise
    return dict(id=doc_id, duplicate=False, chunks=len(chunks), provider=provider())


def documents():
    with closing(connect()) as db:
        return [
            dict(
                id=r[0],
                name=r[1],
                category=r[2],
                provider=r[3],
                locations=r[4],
                chunks=r[5],
            )
            for r in db.execute(
                "SELECT d.id,d.name,d.category,d.provider,d.pages,COUNT(c.id) FROM documents d LEFT JOIN chunks c ON c.document_id=d.id GROUP BY d.id ORDER BY d.rowid DESC"
            )
        ]


def search(question, document_ids=(), limit=4, allow_external=False):
    available = {d["id"] for d in documents() if d["provider"] == index_identity()}
    if document_ids and not set(document_ids) <= available:
        raise ValueError("Unknown document or incompatible embedding provider.")
    if not available:
        return []
    col = collection()
    count = col.count()
    if not count:
        return []
    where = {"document_id": {"$in": list(document_ids)}} if document_ids else None
    result = col.query(
        query_embeddings=embeddings([question], allow_external),
        n_results=min(limit, count),
        where=where,
    )
    answer = []
    with closing(connect()) as db:
        for cid, distance in zip(result["ids"][0], result["distances"][0]):
            if distance > 0.88:
                continue
            row = db.execute(
                "SELECT c.document_id,d.name,c.location,c.text FROM chunks c JOIN documents d ON d.id=c.document_id WHERE c.id=?",
                (cid,),
            ).fetchone()
            if row:
                if provider() == "local-hash-v1":
                    ignored = {
                        "what",
                        "does",
                        "the",
                        "say",
                        "about",
                        "how",
                        "is",
                        "a",
                        "to",
                        "of",
                        "and",
                        "in",
                        "for",
                        "sop",
                        "please",
                        "can",
                        "i",
                        "should",
                    }
                    terms = set(re.findall(r"[a-z0-9]+", question.lower())) - ignored
                    chinese = set(re.findall(r"[\u4e00-\u9fff]{2}", question))
                    if not (
                        terms & set(re.findall(r"[a-z0-9]+", row[3].lower()))
                    ) and not any(t in row[3] for t in chinese):
                        continue
                answer.append(
                    dict(
                        chunk_id=cid,
                        document_id=row[0],
                        name=row[1],
                        location=row[2],
                        quote=row[3],
                        similarity=round(1 - distance, 3),
                    )
                )
    return answer


def reindex(doc_id, allow_external=False):
    with closing(connect()) as db:
        row = db.execute(
            "SELECT provider FROM documents WHERE id=?", (doc_id,)
        ).fetchone()
        if not row:
            raise ValueError("Document not found.")
        if row[0] != index_identity():
            raise ValueError(
                "Provider changed. Use a new document index; do not mix vector spaces."
            )
        chunks = db.execute(
            "SELECT id,text FROM chunks WHERE document_id=? ORDER BY rowid", (doc_id,)
        ).fetchall()
        collection().upsert(
            ids=[c[0] for c in chunks],
            embeddings=embeddings([c[1] for c in chunks], allow_external),
            metadatas=[{"document_id": doc_id} for _ in chunks],
        )
    return dict(id=doc_id, chunks=len(chunks), reindexed=True)
